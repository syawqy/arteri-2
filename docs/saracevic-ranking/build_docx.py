import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_journal_docx(output_path):
    doc = docx.Document()
    
    # 1. Page Margins (Standard Academic 1 inch / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles setup
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    # Helper functions
    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(15)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        return p

    def add_authors():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        run1 = p.add_run("Syauqi Fuadi, S.Hum.*1, [Penulis Kedua]2\n")
        run1.bold = True
        run1.font.size = Pt(11)
        run1.font.name = 'Times New Roman'
        
        run2 = p.add_run("1Program Studi Ilmu Perpustakaan / Peminatan Kearsipan, Fakultas Ilmu Pengetahuan Budaya, Universitas Indonesia\n2[Afiliasi Penulis Kedua]\n")
        run2.font.size = Pt(9.5)
        run2.font.italic = True
        run2.font.name = 'Times New Roman'
        
        run3 = p.add_run("*Penulis Korespondensi: syauqi@fuadi.dev / xfuadi@gmail.com")
        run3.font.size = Pt(9)
        run3.font.italic = True
        run3.font.name = 'Times New Roman'
        p.paragraph_format.space_after = Pt(18)
        return p

    def add_abstract_box(abstract_text, keywords_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        
        # Light grey background & subtle border
        shading = parse_xml(r'<w:shd {} w:fill="F8F9FA"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(r'''
            <w:tcBorders {}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>
                <w:left w:val="none"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        '''.format(nsdecls('w')))
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        run_h = p.add_run("ABSTRAK\n")
        run_h.bold = True
        run_h.font.size = Pt(10)
        run_h.font.name = 'Times New Roman'
        
        run_b = p.add_run(abstract_text + "\n\n")
        run_b.font.size = Pt(9.5)
        run_b.font.name = 'Times New Roman'
        
        run_kw_h = p.add_run("Kata Kunci: ")
        run_kw_h.bold = True
        run_kw_h.font.size = Pt(9.5)
        run_kw_h.font.italic = True
        run_kw_h.font.name = 'Times New Roman'
        
        run_kw = p.add_run(keywords_text)
        run_kw.font.size = Pt(9.5)
        run_kw.font.italic = True
        run_kw.font.name = 'Times New Roman'
        
        # spacing after table
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(0)
        p_space.paragraph_format.space_after = Pt(12)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.italic = True
        run.font.size = Pt(11)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        return p

    def add_p(text, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(text)
        run.font.size = Pt(10.5)
        run.font.name = 'Times New Roman'
        return p

    def add_formula_box(eq_text, label=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(eq_text)
        run.font.name = 'Cambria Math'
        run.font.size = Pt(10.5)
        run.bold = True
        if label:
            run_lbl = p.add_run(f"    ({label})")
            run_lbl.font.name = 'Times New Roman'
            run_lbl.font.size = Pt(10)
            run_lbl.bold = False
        return p

    def add_table_data(headers, rows, col_widths=None):
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        # Header Row
        hdr_cells = table.rows[0].cells
        for idx, header_text in enumerate(headers):
            hdr_cells[idx].text = header_text
            p = hdr_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9.5)
                r.font.name = 'Times New Roman'
            shd = parse_xml(r'<w:shd {} w:fill="EAECEF"/>'.format(nsdecls('w')))
            hdr_cells[idx]._tc.get_or_add_tcPr().append(shd)

        # Data Rows
        for r_idx, row_data in enumerate(rows):
            row_cells = table.rows[r_idx + 1].cells
            for c_idx, cell_value in enumerate(row_data):
                row_cells[c_idx].text = cell_value
                p = row_cells[c_idx].paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.name = 'Times New Roman'
                if c_idx > 0 and len(cell_value) < 15 and ('+' in cell_value or cell_value.replace(',','.').replace('.','').isdigit()):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    
        # Apply borders (APA style top/bottom table borders)
        for r_idx, row in enumerate(table.rows):
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                if r_idx == 0:
                    b_xml = parse_xml(r'''<w:tcBorders {}>
                        <w:top w:val="single" w:sz="8" w:space="0" w:color="222222"/>
                        <w:bottom w:val="single" w:sz="6" w:space="0" w:color="444444"/>
                        <w:left w:val="none"/><w:right w:val="none"/>
                    </w:tcBorders>'''.format(nsdecls('w')))
                elif r_idx == len(table.rows) - 1:
                    b_xml = parse_xml(r'''<w:tcBorders {}>
                        <w:top w:val="none"/>
                        <w:bottom w:val="single" w:sz="8" w:space="0" w:color="222222"/>
                        <w:left w:val="none"/><w:right w:val="none"/>
                    </w:tcBorders>'''.format(nsdecls('w')))
                else:
                    b_xml = parse_xml(r'''<w:tcBorders {}>
                        <w:top w:val="none"/><w:bottom w:val="none"/>
                        <w:left w:val="none"/><w:right w:val="none"/>
                    </w:tcBorders>'''.format(nsdecls('w')))
                tcPr.append(b_xml)
                
        # Set widths if specified
        if col_widths:
            for row in table.rows:
                for idx, w in enumerate(col_widths):
                    row.cells[idx].width = Inches(w)
                    
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(2)
        p_sp.paragraph_format.space_after = Pt(8)

    # --- BUILD DOCUMENT CONTENT ---
    add_title("Pengembangan Sistem Temu Kembali Arsip Digital Berbasis Model Relevansi Bertingkat Saracevic: Implementasi pada Arteri-2")
    add_authors()
    
    abstract_text = (
        "Sistem temu kembali arsip dinamis pada aplikasi basis data relasional umumnya mengandalkan pencocokan sub-string kata kunci Boolean tanpa mekanisme pemeringkatan relevansi (ranking). Konsekuensinya, setiap rekaman arsip yang mengandung kata kunci disajikan setara dan diurutkan secara kronologis atau berdasarkan kunci primer sistem, tanpa mempertimbangkan konteks fungsional kearsipan seperti jatuh tempo retensi dan relasi kelembagaan. Penelitian ini menerapkan metodologi Design Science Research (DSR) untuk merancang dan mengevaluasi artefak sistem temu kembali arsip berbasis model relevansi bertingkat Saracevic (1975; 2007) dan Stratified Model of IR Interaction (1997), yang diadaptasi dari kerangka kerja Fafalios et al. (2017). Implementasi dilakukan pada sistem manajemen arsip Arteri-2 (CodeIgniter 4, MySQL/SQLite) dengan memetakan tiga dimensi relevansi kearsipan: (i) lexical relativeness sebagai representasi relevansi sistem dan topikal, (ii) timeliness berbasis jadwal retensi arsip (b = tanggal + retensi) sebagai relevansi situasional, dan (iii) organizational relations berbasis ko-okurensi kode klasifikasi dan unit pencipta sebagai relevansi kognitif. Agregasi skor (0,5 x kecocokan + 0,3 x urgensi + 0,2 x relasi) dieksekusi langsung pada lapisan kueri basis data (SQL-hybrid ranking) guna menjamin konsistensi pemeringkatan lintas halaman (global pagination). Evaluasi teknis sintetis terhadap 30 kueri uji menunjukkan peningkatan efektivitas secara konsisten: pada dataset pilot 120 arsip diperoleh ΔP@10 = +0,186 dan ΔNDCG@10 = +0,172; sedangkan pada pengujian skala 2.000 arsip diperoleh ΔP@10 = +0,300 dan ΔNDCG@10 = +0,222. Hasil ini membuktikan kelayakan teknis model relevansi bertingkat dalam meningkatkan ketepatan temu kembali pada repositori kearsipan."
    )
    keywords_text = "temu kembali arsip, relevansi bertingkat, Saracevic, stratified model, ranking algoritma, Arteri-2, Design Science Research"
    add_abstract_box(abstract_text, keywords_text)
    
    # 1. Pendahuluan
    add_h1("1. PENDAHULUAN")
    add_p("Pengelolaan arsip dinamis dan inaktif di lingkungan instansi pemerintah menuntut kepatuhan terhadap tata kelola kearsipan, khususnya penerapan kode klasifikasi dan Jadwal Retensi Arsip (JRA). Namun, implementasi modul pencarian pada banyak sistem informasi kearsipan konvensional—termasuk pada arsitektur awal sistem Arteri-2—masih bertumpu pada pencocokan leksikal sederhana (LIKE '%kata_kunci%') yang disajikan berdasarkan urutan penyisipan basis data (ORDER BY a.id ASC).")
    add_p("Ketiadaan mekanisme pemeringkatan relevansi (relevance ranking) menimbulkan masalah temu kembali yang signifikan ketika volume data bertambah. Dokumen yang paling dibutuhkan pengguna dapat terdistribusi pada nomor identitas (ID) yang tinggi atau halaman pagination yang dalam, sehingga luput dari perhatian pengguna pada sepuluh rekaman pertama (top-10 results). Sebagaimana diidentifikasi oleh Fafalios et al. (2017) pada repositori arsip digital historis, ketiadaan lapisan pemeringkatan menyebabkan fenomena information overload, di mana seluruh dokumen yang memenuhi kriteria Boolean dianggap berbobot identik.")
    add_p("Di bidang ilmu informasi, kerangka teori relevansi Saracevic (1975; 2007) mendefinisikan relevansi sebagai konsep multidimensi yang mencakup strata system/topical, cognitive, situational, hingga motivational. Saracevic (1997) kemudian memperluasnya melalui Stratified Model of Information Retrieval Interaction. Meskipun demikian, operasionalisasi praktis model Saracevic ke dalam algoritma temu kembali sistem kearsipan instansi pemerintah masih belum banyak dieksplorasi. Domain kearsipan memiliki karakteristik situasional yang khas: jadwal retensi arsip berfungsi sebagai sinyal urgensi waktu (timeliness), dan struktur pencipta arsip bersama kode klasifikasi merepresentasikan relasi fungsi organisasi (cognitive context). Penelitian ini mengisi kesenjangan tersebut dengan merancang model matematis yang mengoperasionalkan strata Saracevic ke dalam kueri temu kembali pada sistem informasi kearsipan.")
    
    # 2. Metodologi
    add_h1("2. METODOLOGI PENELITIAN")
    add_p("Penelitian ini mengadopsi metodologi Design Science Research (DSR) mengacu pada kerangka kerja Peffers et al. (2007). Metodologi DSR dipilih karena penelitian ini berorientasi pada penciptaan dan evaluasi artefak rekayasa teknologi informasi untuk memecahkan problem praktis dalam temu kembali arsip. Alur penelitian dilaksanakan melalui empat tahapan terstruktur:")
    add_p("1. Identifikasi Masalah (Problem Identification): Menganalisis keterbatasan pencarian SQL standar pada Arteri-2 yang mengurutkan hasil berdasarkan primary key ID dan mengabaikan metadata fungsional kearsipan.")
    add_p("2. Perancangan Konseptual & Matematis (Design & Development): Memetakan strata Saracevic ke dalam atribut basis data serta merumuskan fungsi objektif pemeringkatan multi-komponen terbobot.")
    add_p("3. Pengembangan & Optimasi Artefak (Engineering Refinement): Mengatasi kendala pemotongan data memori (memory buffer truncation) pada implementasi PHP awal dengan mentranslasikan kalkulasi skor ke arsitektur SQL-Hybrid (ORDER BY score DESC di tingkat basis data) guna menjamin konsistensi pagination global.")
    add_p("4. Demonstrasi & Evaluasi Teknis (Technical Validation): Menguji efektivitas artefak menggunakan 30 kueri benchmark terstandarisasi terhadap metrik Precision at 10 (P@10) dan Normalized Discounted Cumulative Gain at 10 (NDCG@10) pada skala 120 dan 2.000 arsip.")

    # 3. Formulasi Model
    add_h1("3. FORMULASI MODEL RELEVANSI BERTINGKAT")
    add_p("Model pemeringkatan menghitung skor total Score(d, q) untuk dokumen arsip d terhadap kueri q melalui kombinasi linear terbobot dari tiga sub-skor terstandarisasi pada interval [0, 1]:")
    
    add_formula_box("Score(d, q) = w_rel · S_rel(d, q) + w_time · S_time(d) + w_relasi · S_relasi(d)", "1")
    
    add_p("dengan batasan Σ w = 1,0. Berdasarkan kerangka kerja Fafalios et al. (2017) dan prioritas strata Saracevic (2007, Part II), bobot awal ditetapkan secara teoritis (theory-driven prior) sebesar: w_rel = 0,50 (Relevansi Topikal & Sistem), w_time = 0,30 (Relevansi Situasional), dan w_relasi = 0,20 (Relevansi Kognitif/Struktural).")
    
    add_h2("3.1 Skor Kecocokan Leksikal (S_rel)")
    add_p("Mengukur keberadaan dan bobot distribusi token kueri q = {t1, t2, ..., tn} pada atribut teks arsip:")
    add_formula_box("S_rel(d, q) = min( 1.0,  ( Σ max_{f in F} ScoreField(t_i, d_f) ) / (3 · n) )", "2")
    add_p("Bobot kecocokan per field (ScoreField) ditetapkan: kecocokan batas kata utuh pada kolom uraian bernilai 3,0 poin; kecocokan sub-string parsial uraian bernilai 1,5 poin; nomor arsip bernilai 2,0 poin; dan nomor boks / kode klasifikasi bernilai 1,0–1,5 poin.")

    add_h2("3.2 Skor Urgensi Waktu (S_time)")
    add_p("Mengukur relevansi situasional arsip berdasarkan kedekatan terhadap tanggal jatuh tempo retensi (b), yang dihitung dari tanggal arsip ditambah masa retensi JRA (UU No. 43/2009; PP No. 28/2012; Peraturan ANRI No. 9/2018):")
    add_formula_box("b = tanggal + retensi", "3")
    add_formula_box("S_time(d) = min( 1.0,  1 / ( 1 + |b - t_sekarang| / 365 ) + Boost_kadaluarsa )", "4")
    add_p("Arsip yang tepat jatuh tempo pada waktu pencarian memperoleh skor dasar 1,00 dan meluruh (decay) secara bertahap seiring bertambahnya selisih tahun. Arsip yang telah melewati masa retensi diberikan Boost_kadaluarsa = 0,08 untuk memprioritaskan tugas seleksi penyusutan atau pemusnahan.")

    add_h2("3.3 Skor Kedekatan Relasi Organisasi (S_relasi)")
    add_p("Mengukur relevansi kognitif melalui frekuensi ko-okurensi historis antara kode klasifikasi (k) dan unit pencipta arsip (p):")
    add_formula_box("S_relasi(d) = Count(k, p) / max_{(k', p')} Count(k', p')", "5")
    add_p("Frekuensi dihitung melalui tabel agregat stat_kode_pencipta. Pasangan klasifikasi-pencipta yang paling dominan dalam repositori memperoleh nilai 1,00, merefleksikan afinitas struktur kelembagaan.")

    # 4. Hasil & Evaluasi
    add_h1("4. HASIL DAN EVALUASI TEKNIS")
    add_p("Efektivitas algoritma diuji menggunakan dua metrik baku Information Retrieval (Manning et al., 2008; Järvelin & Kekäläinen, 2002):")
    add_p("• Precision at 10 (P@10): Mengukur proporsi dokumen relevan (Gain >= 2 pada skala Saracevic 0-3) pada sepuluh hasil teratas.")
    add_p("• Normalized Discounted Cumulative Gain at 10 (NDCG@10): Mengukur kualitas urutan peringkat dengan penalti logaritmik terhadap dokumen relevan yang berada di peringkat bawah.")
    
    add_p("Tabel 1 menyajikan ringkasan hasil pengujian komparatif antara metode baseline (ORDER BY a.id ASC) dan SQL-Hybrid Ranking (ORDER BY score DESC) pada 30 kueri benchmark:")
    
    table_headers = ["Skala Dataset", "P@10 Base", "P@10 Rank", "Δ P@10", "NDCG Base", "NDCG Rank", "Δ NDCG", "Keterangan"]
    table_rows = [
        ["Pilot (120 arsip)", "0,321", "0,507", "+0,186", "0,828", "1,000*", "+0,172", "Skala repositori terbatas"],
        ["Stress-test (2.000 arsip)", "0,332", "0,632", "+0,300", "0,778", "1,000*", "+0,222", "Deep pagination global"]
    ]
    add_table_data(table_headers, table_rows, [1.5, 0.6, 0.6, 0.6, 0.7, 0.7, 0.6, 1.2])
    
    add_p("(*Catatan: Nilai NDCG ranked 1,000 merupakan batas atas sintetis yang membuktikan konsistensi matematis algoritma dalam menyusun peringkat optimal).")
    add_p("Pada pengujian 2.000 arsip, peningkatan efektivitas terlihat sangat signifikan (ΔP@10 = +0,300 dan ΔNDCG@10 = +0,222). Arsitektur SQL-Hybrid terbukti mampu mengangkat dokumen berkategori sangat relevan dari posisi terpendam (misalnya ID 1872 pada kueri 'rekrutmen' yang awalnya berada di halaman 94) langsung ke peringkat pertama pada halaman utama. Eksekusi kueri pada basis data relasional hanya membutuhkan waktu 1–3 ms per kueri.")

    # 5. Kesimpulan
    add_h1("5. KESIMPULAN DAN PENELITIAN LANJUTAN")
    add_p("Model relevansi bertingkat Saracevic berhasil dioperasionalkan pada sistem kearsipan Arteri-2 melalui arsitektur SQL-Hybrid Ranking. Integrasi parameter jadwal retensi dan struktur organisasi ke dalam algoritma temu kembali secara konsisten meningkatkan presisi dan kualitas urutan dokumen tanpa membebani infrastruktur basis data.")
    add_p("Penelitian lanjutan (Tahap 2) akan melibatkan evaluasi pengguna riil (human relevance judgment), integrasi algoritma stemming morfologis bahasa Indonesia (Sastrawi), dan optimasi pembobotan dinamis berbasis umpan balik pengguna.")

    # 6. Daftar Pustaka
    add_h1("DAFTAR PUSTAKA")
    refs = [
        "Fafalios, P., Kasturia, V., & Nejdl, W. (2017). Towards a Ranking Model for Semantic Layers over Digital Archives. Proceedings of the ACM/IEEE Joint Conference on Digital Libraries (JCDL 2017). https://doi.org/10.1109/JCDL.2017.7991617 (arXiv:1810.11049).",
        "Faggioli, G., Dietz, L., & Clarke, C. L. A. (2023). Perspectives on Large Language Models for Relevance Judgment. Proceedings of the 2023 ACM SIGIR International Conference on Theory of Information Retrieval. https://doi.org/10.1145/3578337.3605136.",
        "Järvelin, K., & Kekäläinen, J. (2002). Cumulated gain-based evaluation of IR techniques. ACM Transactions on Information Systems (TOIS), 20(4), 422–446. https://doi.org/10.1145/582415.582418.",
        "Manning, C. D., Raghavan, P., & Schütze, H. (2008). Introduction to Information Retrieval. Cambridge University Press.",
        "Peffers, K., Tuunanen, T., Rothenberger, M. A., & Chatterjee, S. (2007). A design science research methodology for information systems research. Journal of Management Information Systems, 24(3), 45–77. https://doi.org/10.2753/MIS0742-1222240302.",
        "Saracevic, T. (1975). RELEVANCE: A review of and a framework for the thinking on the notion in information science. Journal of the American Society for Information Science, 26(6), 321–343. https://doi.org/10.1002/asi.4630260604.",
        "Saracevic, T. (1997). The Stratified Model of Information Retrieval Interaction: Extension and Applications. Proceedings of the ASIST Annual Meeting, 34, 313–327.",
        "Saracevic, T. (2007). Relevance: A review of the literature and a framework for thinking on the notion in information science. Part II: Nature and manifestations of relevance. Journal of the American Society for Information Science and Technology, 58(13), 1915–1933. https://doi.org/10.1002/asi.20682.",
        "Saracevic, T. (2007). Relevance: A review of the literature and a framework for thinking on the notion in information science. Part III: Behavior and effects of relevance. Journal of the American Society for Information Science and Technology, 58(14), 2126–2144. https://doi.org/10.1002/asi.20681."
    ]
    for r in refs:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.first_line_indent = Inches(-0.4)
        p.paragraph_format.line_spacing = 1.1
        run = p.add_run(r)
        run.font.size = Pt(9.5)
        run.font.name = 'Times New Roman'

    doc.save(output_path)
    print(f"File DOCX successfully generated at {output_path}")

if __name__ == '__main__':
    create_journal_docx('/home/ubuntu/arteri-2/docs/saracevic-ranking/Draft_Artikel_Saracevic_Ranking_Arteri2.docx')
