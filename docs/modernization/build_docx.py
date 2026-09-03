import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_modernization_journal_docx(output_path):
    doc = docx.Document()
    
    # 1. Page Margins (Standard Academic 1 inch / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles setup
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(15)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        return p

    def add_authors():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        run1 = p.add_run("Syauqi Fuadi, S.Hum.*1, [Penulis Kedua]2\n")
        run1.bold = True
        run1.font.size = Pt(11)
        run1.font.name = 'Times New Roman'
        
        run2 = p.add_run("1Program Studi Ilmu Perpustakaan / Peminatan Kearsipan, Fakultas Ilmu Pengetahuan Budaya, Universitas Indonesia\n2[Afiliasi Penulis Kedua]\n")
        run2.font.size = Pt(9.5)
        run2.font.italic = True
        run2.font.name = 'Times New Roman'
        
        run3 = p.add_run("*Penulis Korespondensi: syauqi@fuadi.dev / xfuadi@gmail.com")
        run3.font.size = Pt(9)
        run3.font.italic = True
        run3.font.name = 'Times New Roman'
        p.paragraph_format.space_after = Pt(18)
        return p

    def add_abstract_box(abstract_text, keywords_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        
        shading = parse_xml(r'<w:shd {} w:fill="F8F9FA"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(r'''
            <w:tcBorders {}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>
                <w:left w:val="none"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        '''.format(nsdecls('w')))
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        run_h = p.add_run("ABSTRAK\n")
        run_h.bold = True
        run_h.font.size = Pt(10)
        run_h.font.name = 'Times New Roman'
        
        run_b = p.add_run(abstract_text + "\n\n")
        run_b.font.size = Pt(9.5)
        run_b.font.name = 'Times New Roman'
        
        run_kw_h = p.add_run("Kata Kunci: ")
        run_kw_h.bold = True
        run_kw_h.font.size = Pt(9.5)
        run_kw_h.font.italic = True
        run_kw_h.font.name = 'Times New Roman'
        
        run_kw = p.add_run(keywords_text)
        run_kw.font.size = Pt(9.5)
        run_kw.font.italic = True
        run_kw.font.name = 'Times New Roman'
        
        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        return p

    def add_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        return p

    def add_code_block(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        
        shading = parse_xml(r'<w:shd {} w:fill="F4F6F8"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(r'''
            <w:tcBorders {}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="E1E4E8"/>
                <w:left w:val="single" w:sz="12" w:space="0" w:color="0366D6"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E1E4E8"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="E1E4E8"/>
            </w:tcBorders>
        '''.format(nsdecls('w')))
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(code_text)
        run.font.name = 'Courier New'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- DOCUMENT GENERATION ---
    add_title("Modernisasi dan Rekayasa Ulang Sistem Manajemen Arsip Digital Berbasis Web: Dari Arteri-1 Menuju Arsitektur Arteri-2")
    add_authors()
    
    abstrak_text = (
        "Sistem informasi kearsipan instansi pemerintah yang dikembangkan pada era awal web umumnya bertumpu pada arsitektur monolitik PHP legacy (seperti CodeIgniter 3 atau PHP prosedural). Seiring berjalannya waktu, sistem warisan (legacy systems) tersebut mengalami degradasi kualitas perangkat lunak (software aging), tingginya akumulasi technical debt, kerentanan keamanan (seperti injeksi SQL, ketiadaan perlindungan CSRF, dan penyimpanan kata sandi tidak aman), serta inkompatibilitas terhadap runtime modern PHP 8.x. Artikel ini menyajikan studi kasus rekayasa ulang perangkat lunak (software re-engineering) dan modernisasi sistem manajemen arsip digital open-source dari Arteri-1 (CodeIgniter 3) menuju Arteri-2 (CodeIgniter 4) menggunakan metodologi Design Science Research (DSR) yang diintegrasikan dengan taksonomi re-engineering Chikofsky dan Cross (1990).\n\n"
        "Proses rekayasa ulang mencakup tiga tahapan inti: (i) reverse engineering dan audit kerentanan sistem warisan, (ii) restructuring arsitektur menuju model MVC terstruktur dengan autoloading PSR-4 dan strict-typing PHP 8.4, serta (iii) forward engineering melalui penambahan lapisan keamanan modern (mitigasi OWASP Top 10, autentikasi berbasis Bcrypt, role-based access control, dan pencatatan jejak audit komprehensif) serta penyediaan RESTful API v1 terstandar OpenAPI 3.0. Evaluasi kualitas perangkat lunak mengacu pada standar ISO/IEC 25010 menunjukkan peningkatan signifikan pada karakteristik Maintainability, Security, dan Compatibility. Pengujian regresi otomatis (36 berkas uji PHPUnit dengan status 100% passing) dan audit keamanan membuktikan bahwa Arteri-2 berhasil mentransformasikan aplikasi kearsipan warisan menjadi platform modern yang tangguh, aman, dan siap terinteroperabilitas dengan ekosistem Sistem Pemerintahan Berbasis Elektronik (SPBE)."
    )
    keywords_text = "rekayasa ulang perangkat lunak, modernisasi sistem warisan, CodeIgniter 4, manajemen arsip digital, Arteri-2, ISO/IEC 25010, OWASP Top 10, Design Science Research"
    add_abstract_box(abstrak_text, keywords_text)

    # Bagian 1
    add_h1("Bagian 1. Pendahuluan")
    add_h2("Bagian 1.1 Latar Belakang")
    add_p("Pengelolaan arsip dinamis dan inaktif di lingkungan sektor publik menuntut keandalan sistem informasi kearsipan yang mampu menjamin integritas, keotentikan, kerahasiaan, dan keteraksesan data dalam jangka panjang. Pada tahun 2017, aplikasi sistem informasi kearsipan berbasis web sumber terbuka Arteri generasi pertama (Arteri-1) dirilis (https://github.com/dicarve/arteri). Arteri-1 dirancang sebagai perangkat lunak pencatatan, penataan, dan temu kembali berkas arsip yang dibangun menggunakan kerangka kerja CodeIgniter 3 dengan lingkungan runtime PHP 5.6 hingga PHP 7.x.")
    add_p("Kendati Arteri-1 berhasil memfasilitasi kebutuhan operasional dasar kearsipan dan sirkulasi peminjaman pada masa awal perilisannya, sistem tersebut memiliki berbagai kekurangan mendasar yang menjadikannya tidak lagi memadai untuk lanskap operasional dan keamanan saat ini. Seiring berjalannya waktu, akumulasi kekurangan teknis (technical debt) dan kerentanan keamanan (security vulnerabilities) pada Arteri-1 menjadi faktor pendorong utama dilakukannya rekayasa ulang menuju Arteri-2:\\n"
          "1. Kerentanan Keamanan Kritis: Implementasi pengamanan pada Arteri-1 belum memenuhi standar keamanan web modern. Hal ini mencakup penggunaan algoritma hashing kata sandi usang (md5/sha1 tanpa salt yang kuat), proteksi CSRF yang tidak aktif secara default, penanganan kueri basis data parsial yang belum sepenuhnya menerapkan parameterized prepared statements (meningkatkan risiko SQL Injection), serta keterbatasan kontrol akses langsung terhadap berkas arsip digital.\\n"
          "2. Ketiadaan Modul Jejak Audit (Audit Trail): Arteri-1 tidak memiliki modul pencatatan jejak aktivitas pengguna (system audit logging), padahal akuntabilitas dan pencatatan riwayat akses/modifikasi dokumen merupakan mandat kepatuhan standar tata kelola kearsipan (ISO 15489).\\n"
          "3. Keterbatasan Ekosistem dan Interoperabilitas: Arteri-1 dibangun sebagai monolit tertutup tanpa dukungan antarmuka pemrograman aplikasi (RESTful API), sehingga menyulitkan integrasi data dengan ekosistem aplikasi lain di lingkungan instansi pemerintah (SPBE).\\n"
          "4. Keusangan Tumpukan Teknologi (Technology Obsolescence): Berakhirnya masa dukungan resmi (End-of-Life / EOL) pada runtime PHP 5.6 dan PHP 7.x serta penghentian pengembangan aktif kerangka kerja CodeIgniter 3 menimbulkan risiko keamanan server dan inkompatibilitas fatal saat dijalankan pada infrastruktur server modern berbasis PHP 8.x.\\n"
          "5. Ketiadaan Pengujian Otomatis (Automated Test Suite): Arteri-1 tidak dibekali mekanisme pengujian unit terotomatisasi, sehingga meningkatkan risiko regresi fungsi saat kode dipelihara.\\n\\n"
          "Faktor-faktor kelemahan fungsional, keamanan, dan keusangan tumpukan teknologi inilah yang mendasari urgensi pengembangan Arteri-2 melalui pendekatan rekayasa ulang perangkat lunak (software re-engineering).")

    add_h2("Bagian 1.2 Kesenjangan Masalah dan Motivasi Re-engineering")
    add_p("Sesuai dengan Hukum Evolusi Perangkat Lunak Lehman (Lehman's Laws of Software Evolution), suatu sistem perangkat lunak yang beroperasi di lingkungan dunia nyata harus terus diadaptasikan secara berkelanjutan; jika tidak, sistem tersebut akan mengalami degradasi kualitas dan utilitas yang semakin tinggi (Lehman, 1996). Alih-alih membangun sistem dari nol (scratch rewrite) yang berisiko membuang logika domain kearsipan yang telah terbukti fungsional sejak rilis 2017, pendekatan software re-engineering dipilih sebagai strategi terstruktur, hemat biaya, dan minim risiko regresi untuk mentransformasikan Arteri-1 menjadi Arteri-2 dengan tumpukan teknologi CodeIgniter 4 dan PHP 8.x.")

    add_h2("Bagian 1.3 Pertanyaan Penelitian")
    add_p("Penelitian ini bertujuan mendokumentasikan, merumuskan arsitektur, dan mengevaluasi proses rekayasa ulang sistem kearsipan Arteri dari generasi pertama ke generasi kedua (Arteri-2). Pertanyaan penelitian yang diajukan adalah:\n"
          "• RQ1. Bagaimana identifikasi technical debt dan pemetaan kerentanan keamanan pada arsitektur sistem kearsipan warisan Arteri-1?\n"
          "• RQ2. Bagaimana rancangan arsitektur rekayasa ulang menuju CodeIgniter 4 (PHP 8.4) yang mempertahankan integritas proses bisnis kearsipan sekaligus mengadopsi prinsip desain modern?\n"
          "• RQ3. Sejauh mana implementasi Arteri-2 meningkatkan kualitas perangkat lunak ditinjau dari karakteristik Security, Maintainability, dan Interoperability berbasis standar ISO/IEC 25010 dan OWASP Top 10?")

    # Bagian 2
    add_h1("Bagian 2. Landasan Teori")
    add_h2("Bagian 2.1 Taksonomi Rekayasa Ulang Perangkat Lunak (Chikofsky & Cross, 1990)")
    add_p("Chikofsky dan Cross (1990) mendefinisikan software re-engineering sebagai proses pemeriksaan dan pengubahan suatu sistem perangkat lunak untuk merekonstruksinya ke dalam bentuk baru serta mengimplementasikan bentuk baru tersebut. Kerangka ini membedakan tiga aktivitas utama: (i) Reverse Engineering (analisis sistem sasaran); (ii) Restructuring (transformasi struktur internal kode); dan (iii) Forward Engineering (rekayasa maju dengan penambahan fitur dan pengerasan keamanan).")

    add_h2("Bagian 2.2 Model Kualitas Perangkat Lunak ISO/IEC 25010")
    add_p("Standar ISO/IEC 25010 menyediakan model evaluasi kualitas produk perangkat lunak. Evaluasi modernisasi ini difokuskan pada karakteristik Maintainability (modularitas, keterujian), Security (kerahasiaan, integritas, akuntabilitas), dan Compatibility / Interoperability.")

    # Bagian 3
    add_h1("Bagian 3. Metodologi Penelitian")
    add_p("Penelitian ini mengadopsi metode Design Science Research (DSR) (Peffers et al., 2007) yang dipadukan dengan siklus Re-engineering Chikofsky dan Cross (1990) dalam empat tahapan operasional:")
    
    img_path = "/home/ubuntu/arteri-2/docs/modernization/images/dsr_methodology_flow.png"
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run("Gambar 1. Kerangka Metodologi Rekayasa Ulang Arteri (DSR + Re-engineering Lifecycle)")
        run_cap.font.italic = True
        run_cap.font.size = Pt(9.5)
        run_cap.font.name = 'Times New Roman'
        
    add_p("Empat tahapan terstruktur tersebut mencakup:\\n"
          "1. Identifikasi Masalah & Audit Sistem Warisan (Reverse Engineering): Melakukan audit menyeluruh terhadap kode sumber repositori Arteri-1 (CodeIgniter 3.1.x), mengidentifikasi kerentanan keamanan warisan (injeksi SQL, hashing kata sandi MD5/SHA1, ketiadaan proteksi CSRF), merekonstruksi model relasional basis data (data_arsip, master_kode, sirkulasi), serta memetakan technical debt dan dependensi PHP EOL.\\n"
          "2. Perancangan Arsitektur Target & Spesifikasi Solusi (Restructuring): Merancang ulang struktur direktori aplikasi dengan memisahkan document root publik (public/index.php) dari logika inti aplikasi (app/), memigrasikan arsitektur ke CodeIgniter 4 berbasis autoloading PSR-4, menegakkan deklarasi tipe data ketat (strict typing) pada PHP 8.4, dan menyusun skrip migrasi basis data deklaratif yang kompatibel secara agnostik dengan MySQL maupun SQLite.\\n"
          "3. Rekayasa Maju & Implementasi Fitur Baru (Forward Engineering): Membangun mekanisme mitigasi OWASP Top 10 (enkripsi Bcrypt, Auth Filter, Secure File Serving), modul pencatatan jejak audit (SystemLog), modul pemulihan data terhapus (Soft Deletes / Trash), serta merancang dan mengimplementasikan antarmuka RESTful API v1 terstandar OpenAPI 3.0 dengan pengamanan API Key berbasis SHA-256 dan pembatasan laju kueri (rate limiting).\\n"
          "4. Demonstrasi & Evaluasi Kualitas (Evaluation): Mengeksekusi rangkaian pengujian regresi otomatis (test suite) berbasis PHPUnit 11, mengevaluasi karakteristik kualitas sistem mengacu pada standar ISO/IEC 25010 (Maintainability, Security, Compatibility, Functional Suitability, dan Interoperability), serta memverifikasi kepatuhan pengerasan keamanan terhadap matriks risiko OWASP Top 10.")

    # Bagian 4
    add_h1("Bagian 4. Proses Rekayasa Ulang dan Arsitektur Arteri-2")
    add_h2("Bagian 4.1 Dekonstruksi Sistem Warisan (Arteri-1)")
    add_p("Analisis terhadap repositori Arteri-1 (https://github.com/dicarve/arteri) mengidentifikasi sejumlah kelemahan struktural pada penanganan kueri basis data mentah, hash kata sandi usang, dan ketiadaan isolasi berkas dokumen.")

    add_h2("Bagian 4.2 Restrukturisasi Arsitektur Menuju CodeIgniter 4 (Arteri-2)")
    
    # Table 1
    t1 = doc.add_table(rows=10, cols=3)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Aspek Arsitektur", "Arteri-1 (Sistem Warisan)", "Arteri-2 (Hasil Re-engineering)"]
    for i, h in enumerate(headers):
        cell = t1.cell(0, i)
        cell.paragraphs[0].add_run(h).bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shading = parse_xml(r'<w:shd {} w:fill="EAECEF"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)

    rows_t1 = [
        ("Kerangka Kerja (Framework)", "CodeIgniter 3.1.x", "CodeIgniter 4.5.x"),
        ("Versi Runtime Bahasa", "PHP 5.6 / 7.2 (EOL)", "PHP 8.2 / 8.3 / 8.4 (Modern)"),
        ("Enkripsi Kata Sandi", "Hash MD5 / SHA-1", "Bcrypt / Argon2ID via password_hash()"),
        ("Pencegahan Injeksi SQL", "Manual / Parsial", "Query Builder Parameterized PDO Penuh"),
        ("Proteksi CSRF", "Nonaktif Default", "Token CSRF Aktif Global pada POST"),
        ("Jejak Audit (Audit Trail)", "Tidak Ada", "Modul SystemLog mencatat aktivitas"),
        ("Penghapusan Data", "Hard Delete", "Soft Deletes dengan fitur Recycle Bin"),
        ("Interoperabilitas API", "Tidak Ada", "RESTful API v1, API Key Auth, OpenAPI"),
        ("Pengujian Otomatis", "0% Test Coverage", "PHPUnit Test Suite (36 Test Files)")
    ]
    for r_idx, r_data in enumerate(rows_t1, start=1):
        for c_idx, val in enumerate(r_data):
            t1.cell(r_idx, c_idx).paragraphs[0].add_run(val)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h2("Bagian 4.3 Rekayasa Maju: Pengerasan Keamanan dan REST API")
    add_p("Arteri-2 mengimplementasikan Access Control List (ACL) berbasis kode klasifikasi, pengamanan akses file digital via FileController, pencatatan otomatis transaksi pada system_log, serta penyediaan RESTful API v1 dengan otentikasi API Key SHA-256.")

    # Bagian 5
    add_h1("Bagian 5. Evaluasi Kualitas dan Pembahasan")
    add_h2("Bagian 5.1 Evaluasi Kualitas Berdasarkan ISO/IEC 25010")
    
    # Table 2
    t2 = doc.add_table(rows=6, cols=4)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers2 = ["Karakteristik Kualitas", "Kondisi Arteri-1", "Hasil Evaluasi Arteri-2", "Status Peningkatan"]
    for i, h in enumerate(headers2):
        cell = t2.cell(0, i)
        cell.paragraphs[0].add_run(h).bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shading = parse_xml(r'<w:shd {} w:fill="EAECEF"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)

    rows_t2 = [
        ("Maintainability", "Monolitik, tanpa unit test", "PSR-4, modular services, 36 test files", "Sangat Signifikan"),
        ("Security", "Rawan SQLi/XSS, hash usang", "Kepatuhan OWASP Top 10, Bcrypt, CSRF", "Sangat Signifikan"),
        ("Functional Suitability", "Fitur dasar pencatatan", "Soft Deletes, Trash, Advanced Filters", "Meningkat"),
        ("Compatibility", "PHP 7 lama & MySQL", "PHP 8.4, Dual MySQL & SQLite", "Sangat Signifikan"),
        ("Interoperability", "Tertutup (Web UI Monolitik)", "RESTful API v1 & OpenAPI 3.0", "Sangat Signifikan")
    ]
    for r_idx, r_data in enumerate(rows_t2, start=1):
        for c_idx, val in enumerate(r_data):
            p = t2.cell(r_idx, c_idx).paragraphs[0]
            r = p.add_run(val)
            if c_idx == 3:
                r.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h2("Bagian 5.2 Verifikasi Keamanan OWASP Top 10")
    add_p("Audit otomatis (docs/security/owasp-evidence.md) memverifikasi bahwa Arteri-2 lulus pada kontrol Broken Access Control (A01), Cryptographic Failures (A02), Injection (A03), Identification & Authentication Failures (A07), serta Security Logging Failures (A09).")

    add_h2("Bagian 5.3 Hasil Pengujian Regresi Otomatis")
    add_p("Seluruh 36 berkas uji PHPUnit 11 dieksekusi dengan hasil 100% passing, menjamin stabilitas fungsional sistem.")

    # Bagian 6
    add_h1("Bagian 6. Kesimpulan dan Agenda Riset Lanjutan")
    add_p("Modernisasi dari Arteri-1 ke Arteri-2 membuktikan bahwa rekayasa ulang sistem kearsipan warisan dapat dicapai secara efisien. Transformasi arsitektural ini meletakkan fondasi platform yang tangguh untuk pengembangan inovasi kearsipan tingkat lanjut berikutnya.")

    # Daftar Pustaka
    add_h1("Daftar Pustaka")
    refs = [
        "Chikofsky, E. J., & Cross, J. H. (1990). Reverse engineering and design recovery: A taxonomy. IEEE Software, 7(1), 13–17. https://doi.org/10.1109/52.43044",
        "International Organization for Standardization. (2011). Systems and software engineering — Systems and software Quality Requirements and Evaluation (SQuaRE) — System and software quality models (ISO/IEC 25010:2011). ISO. https://www.iso.org/standard/35733.html",
        "Lehman, M. M. (1996). Laws of software evolution revisited. Lecture Notes in Computer Science, 1149, 108–124. https://doi.org/10.1007/bfb0017737",
        "Open Web Application Security Project. (2021). OWASP Top 10:2021 The Ten Most Critical Web Application Security Risks. OWASP Foundation. https://owasp.org/Top10/",
        "Peffers, K., Tuunanen, T., Rothenberger, M. A., & Chatterjee, S. (2007). A design science research methodology for information systems research. Journal of Management Information Systems, 24(3), 45–77. https://doi.org/10.2753/mis0742-1222240302"
    ]
    for rf in refs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(rf)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)

    # Lampiran
    add_h1("Lampiran A. Potongan Kode Rekayasa Ulang Utama")
    code_a1 = '''// app/Models/ArsipModel.php - Migrasi dari string concatenation ke Parameterized Query Builder
public function searchWithCursor(?int $cursor, int $limit, string $keywords, array $filters = []): array
{
    $builder = $this->builder();
    $builder->select('data_arsip.*, master_kode.nama as nama_kode, master_kode.retensi');
    $builder->join('master_kode', 'master_kode.kode = data_arsip.kode', 'left');

    if ($cursor !== null && $cursor > 0) {
        $builder->where('data_arsip.id <', $cursor);
    }

    if (!empty($keywords)) {
        $builder->groupStart()
            ->like('data_arsip.uraian', $keywords)
            ->orLike('data_arsip.noarsip', $keywords)
            ->groupEnd();
    }

    $builder->orderBy('data_arsip.id', 'DESC');
    return $builder->get($limit)->getResultArray();
}'''
    add_code_block(code_a1)

    doc.save(output_path)
    print(f"File DOCX Prequel berhasil dibuat di: {output_path}")

if __name__ == '__main__':
    create_modernization_journal_docx('/home/ubuntu/arteri-2/docs/modernization/Draft_Artikel_Modernisasi_Arteri2.docx')
